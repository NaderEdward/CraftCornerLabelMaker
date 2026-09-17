from __future__ import annotations

import csv
import json
import webbrowser
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from core import paths as core_paths
from core.units import px_to_mm

_TEMPLATE_PATH = Path(__file__).parent / "templates" / "report_template.html"

_PRODUCT_COLOURS = [
    "var(--vinyl-1)", "var(--vinyl-2)", "var(--vinyl-3)",
    "var(--vinyl-4)", "var(--vinyl-5)",
]


def schematics_from_plan(plan: Dict[str, Any], manifest: Dict[str, Any],
                         sheet_numbers: List[int],
                         dpi: int) -> Dict[str, List[Dict[str, Any]]]:
    tile_by_id = {t["tile_id"]: t for t in manifest["tiles"]}
    out: Dict[str, List[Dict[str, Any]]] = {}

    for sheet, number in zip(plan.get("sheets", []), sheet_numbers):
        rects: List[Dict[str, Any]] = []
        for placement in sheet["placements"]:
            meta = tile_by_id.get(placement["tile_id"])
            if meta is None:
                continue
            w_mm, h_mm = meta["w_mm"], meta["h_mm"]
            if placement.get("rotated"):
                w_mm, h_mm = h_mm, w_mm
            rects.append({
                "x": round(px_to_mm(placement["x_px"], dpi), 1),
                "y": round(px_to_mm(placement["y_px"], dpi), 1),
                "w": round(w_mm, 1),
                "h": round(h_mm, 1),
                "p": _slugify(meta.get("region_label", "")),
            })
        out[f"{number:03d}"] = rects

    return out


def _read_xlsx_rows(xlsx_path: Path) -> tuple:
    if not xlsx_path.exists():
        return [], 0
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb["Production"] if "Production" in wb.sheetnames else None
        if ws is None:
            return [], 0
        rows, skipped = [], 0
        headers = None
        from report.csv_log import HEADERS as _HEADERS
        for i, xl_row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(c) if c is not None else "" for c in xl_row]
                continue
            if not xl_row or xl_row[0] is None:
                skipped += 1
                continue
            row_dict = {h: (str(v) if v is not None else "") for h, v in zip(headers, xl_row)}
            if not row_dict.get("Sheet #"):
                skipped += 1
                continue
            rows.append(row_dict)
        wb.close()
        return rows, skipped
    except Exception:
        return [], 0


def _slugify(label: str) -> str:
    return "".join(c.lower() if c.isalnum() else "_" for c in label).strip("_")


def _build_products(rows: List[Dict[str, str]]) -> Dict[str, Dict[str, str]]:
    products: Dict[str, Dict[str, str]] = {}
    for row in rows:
        label = row.get("Product", "")
        if not label:
            continue
        key = _slugify(label)
        if key not in products:
            colour = _PRODUCT_COLOURS[len(products) % len(_PRODUCT_COLOURS)]
            products[key] = {"label": label, "colour": colour}
    return products


def _build_sheets(rows: List[Dict[str, str]],
                  tile_schematics: Dict[str, List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    by_sheet: Dict[str, List[Dict[str, str]]] = defaultdict(list)
    for row in rows:
        by_sheet[row.get("Sheet #", "")].append(row)

    sheets = []
    for sheet_no, sheet_rows in sorted(by_sheet.items(), key=lambda kv: kv[0], reverse=True):
        first = sheet_rows[0]
        out_rows = []
        for r in sheet_rows:
            out_rows.append({
                "ord": r.get("Order #", ""),
                "od": r.get("Order Date", ""),
                "cust": r.get("Customer Name", ""),
                "stu": r.get("Student Name", ""),
                "ar": r.get("Student Name (Arabic)", ""),
                "school": r.get("School", ""),
                "grade": r.get("Grade", ""),
                "p": _slugify(r.get("Product", "")),
                "q": int(r.get("Qty on This Sheet", "0") or 0),
            })
        fill_pct_str = first.get("Fill %", "0")
        try:
            fill = float(fill_pct_str.strip("%")) / 100 if "%" in fill_pct_str else float(fill_pct_str)
        except ValueError:
            fill = 0.0
        sheets.append({
            "no": sheet_no,
            "date": first.get("Export Date", ""),
            "fill": round(fill, 4),
            "forced": False,
            "tiles": tile_schematics.get(sheet_no, []),
            "rows": out_rows,
        })
    return sheets


def _pending_from_plan(plan: Optional[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not plan:
        return []
    pending = []
    for order_number in plan.get("held_back_orders", []):
        pending.append({"ord": order_number, "stu": "", "items": "", "days": 0})
    return pending


def generate_report(xlsx_path: Optional[Path] = None,
                    csv_path: Optional[Path] = None,
                    out_path: Optional[Path] = None,
                    latest_plan: Optional[Dict[str, Any]] = None,
                    tile_schematics: Optional[Dict[str, List[Dict[str, Any]]]] = None,
                    sheet_mm: Optional[Dict[str, float]] = None) -> Path:
    xlsx_path = xlsx_path or core_paths.reports_dir() / "sheets_report.xlsx"
    out_path = out_path or core_paths.reports_dir() / "production_report.html"
    tile_schematics = tile_schematics or {}

    rows, skipped = _read_xlsx_rows(xlsx_path)
    products = _build_products(rows)
    sheets = _build_sheets(rows, tile_schematics)
    pending = _pending_from_plan(latest_plan)

    data = {
        "generated": datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M"),
        "sheet_mm": sheet_mm or {"w": 420, "h": 297},
        "sheets": sheets,
        "pending": pending,
        "skipped_rows": skipped,
    }

    template_html = _TEMPLATE_PATH.read_text(encoding="utf-8")
    html = template_html.replace(
        "/*__PRODUCTS_JSON__*/{}/*__END_PRODUCTS_JSON__*/",
        json.dumps(products, ensure_ascii=False),
    ).replace(
        '/*__DATA_JSON__*/{"generated":"","sheet_mm":{"w":420,"h":297},"sheets":[],"pending":[]}/*__END_DATA_JSON__*/',
        json.dumps(data, ensure_ascii=False),
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html, encoding="utf-8")

    write_report_cache(data, products, out_path.parent)
    return out_path


CACHE_FILENAME = "report_data.json"


def write_report_cache(data: Dict[str, Any], products: Dict[str, Any],
                       reports_dir: Optional[Path] = None) -> Path:
    reports_dir = reports_dir or core_paths.reports_dir()
    reports_dir.mkdir(parents=True, exist_ok=True)
    cache_path = reports_dir / CACHE_FILENAME

    payload = {
        "generated": data.get("generated", ""),
        "sheet_mm": data.get("sheet_mm", {"w": 420, "h": 297}),
        "sheets": data.get("sheets", []),
        "pending": data.get("pending", []),
        "skipped_rows": data.get("skipped_rows", 0),
        "products": products,
    }

    tmp = cache_path.with_suffix(cache_path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    tmp.replace(cache_path)
    return cache_path


def read_report_cache(reports_dir: Optional[Path] = None
                      ) -> Optional[Dict[str, Any]]:
    reports_dir = reports_dir or core_paths.reports_dir()
    cache_path = reports_dir / CACHE_FILENAME
    if not cache_path.exists():
        return None
    try:
        return json.loads(cache_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None


def open_report(out_path: Optional[Path] = None) -> None:
    out_path = out_path or core_paths.reports_dir() / "production_report.html"
    webbrowser.open(out_path.resolve().as_uri())
