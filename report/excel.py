from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from report import csv_log

HEADERS = csv_log.HEADERS


class ReportFileLockedError(RuntimeError):
    pass


def _is_locked(xlsx_path: Path) -> bool:
    lock_file = xlsx_path.parent / f"~${xlsx_path.name}"
    if lock_file.exists():
        return True
    if xlsx_path.exists():
        try:
            with open(xlsx_path, "a+b"):
                pass
        except PermissionError:
            return True
    return False


def record_key(order_id: str, student_name: str) -> tuple:
    return (order_id, student_name)


def index_records(records: List[Dict[str, Any]]) -> Dict[tuple, Dict[str, Any]]:
    index: Dict[tuple, Dict[str, Any]] = {}
    for r in records:
        key = record_key(r["order_id"], r["student_name"])
        existing = index.get(key)
        if existing is None:
            index[key] = dict(r)
            continue
        for field, value in r.items():
            if value and not existing.get(field):
                existing[field] = value
    return index


def build_report_rows(sheet_number: int, export_date: str, plan_sheet: Dict[str, Any],
                      manifest: Dict[str, Any], records_by_key: Dict[tuple, Any],
                      run_id: str) -> List[Dict[str, Any]]:
    tile_by_id = {t["tile_id"]: t for t in manifest["tiles"]}

    groups: Dict[tuple, Dict[str, Any]] = {}
    for placement in plan_sheet["placements"]:
        meta = tile_by_id[placement["tile_id"]]
        key = (meta["order_id"], meta["student_name"], meta["region_id"])
        g = groups.setdefault(key, {"meta": meta, "qty": 0, "tile_ids": []})
        g["qty"] += 1
        g["tile_ids"].append(meta["tile_id"])

    rows: List[Dict[str, Any]] = []
    fill_pct = plan_sheet["utilisation"]
    for (order_id, student_name, region_id), g in groups.items():
        meta = g["meta"]
        record = records_by_key.get(record_key(order_id, student_name), {})
        rows.append({
            "sheet_number": f"{sheet_number:03d}",
            "export_date": export_date,
            "fill_pct": fill_pct,
            "order_number": meta["order_number"],
            "order_date": record.get("order_date", ""),
            "customer_name": meta["customer_name"],
            "student_name": meta["student_name"],
            "student_name_arabic": record.get("student_name_arabic", ""),
            "school": record.get("school", ""),
            "grade": record.get("grade", ""),
            "product": meta["region_label"],
            "qty_on_sheet": g["qty"],
            "template": meta["template_name"],
            "tile_ids": ",".join(sorted(g["tile_ids"])),
            "run_id": run_id,
        })
    return rows


def append_report(xlsx_path: Path, rows: List[Dict[str, Any]]) -> None:
    if _is_locked(xlsx_path):
        raise ReportFileLockedError(
            f"{xlsx_path.name} is open in Excel. Close it and retry, or "
            f"rely on the CSV (already written)."
        )

    if xlsx_path.exists():
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb["Production"] if "Production" in wb.sheetnames else wb.create_sheet("Production")
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Production"

    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        for col, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(name="Arial", bold=True)
        ws.freeze_panes = "A2"

    for r in rows:
        row_vals = csv_log.row_from_report_row(r)
        ws.append([row_vals[h] for h in HEADERS])
        last_row = ws.max_row
        ws.cell(row=last_row, column=HEADERS.index("Fill %") + 1).number_format = "0.0%"
        arabic_col = HEADERS.index("Student Name (Arabic)") + 1
        cell = ws.cell(row=last_row, column=arabic_col)
        cell.alignment = Alignment(horizontal="right")

    for col_idx, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(header) + 2)

    last_col = get_column_letter(len(HEADERS))
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"

    tmp_path = xlsx_path.with_suffix(".xlsx.tmp")
    wb.save(tmp_path)
    tmp_path.replace(xlsx_path)


def rebuild_from_csv(csv_path: Path, xlsx_path: Path) -> None:
    import csv as _csv

    if xlsx_path.exists():
        xlsx_path.unlink()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Production"
    for col, header in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=col, value=header).font = Font(name="Arial", bold=True)
    ws.freeze_panes = "A2"

    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        reader = _csv.DictReader(f)
        for row in reader:
            ws.append([row.get(h, "") for h in HEADERS])

    last_col = get_column_letter(len(HEADERS))
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"

    tmp_path = xlsx_path.with_suffix(".xlsx.tmp")
    wb.save(tmp_path)
    tmp_path.replace(xlsx_path)
