from __future__ import annotations

import json

import pytest

from shopify.review import (
    _flatten_records,
    apply_review,
    diff_review,
    export_review,
)

RECORDS_DICT = {
    "generated_at": "2026-08-04T00:00:00",
    "template_versions": {},
    "warnings": [
        {"order_number": "#1044", "kind": "unmatched_sku",
         "detail": "UNKNOWN-SKU-9", "line_item_id": "1130"},
    ],
    "records": [
        {
            "order_id": "5001", "order_number": "#1042",
            "customer_name": "Sarah Al-Rashidi",
            "student_name": "Ahmed Al-Rashidi",
            "student_name_arabic": "أحمد",
            "school": "Greenfield International", "grade": "Grade 3",
            "telephone": "", "custom_fields": {},
            "template_name": "baby_minnie_corrected", "language": "en",
            "items": [
                {"region_id": "large_labels",   "qty": 2, "sku": "BMINNIE-L"},
                {"region_id": "meeting_labels", "qty": 1, "sku": "BMINNIE-M"},
            ],
        },
        {
            "order_id": "5002", "order_number": "#1045",
            "customer_name": "Yasmine Fouad",
            "student_name": "Laila Fouad",
            "student_name_arabic": "ليلى",
            "school": "", "grade": "KG 2",
            "telephone": "", "custom_fields": {},
            "template_name": "baby_minnie_corrected", "language": "en",
            "items": [
                {"region_id": "large_labels", "qty": 1, "sku": "BMINNIE-L"},
            ],
        },
    ],
}


@pytest.fixture()
def run_dir(tmp_path):
    p = tmp_path / "run"
    p.mkdir()
    (p / "records.json").write_text(
        json.dumps(RECORDS_DICT, ensure_ascii=False), encoding="utf-8"
    )
    return p


def test_export_creates_xlsx(run_dir):
    out = run_dir / "orders_review.xlsx"
    export_review(run_dir / "records.json", out)
    assert out.exists() and out.stat().st_size > 0


def test_export_has_three_sheets(run_dir):
    import openpyxl
    out = run_dir / "orders_review.xlsx"
    export_review(run_dir / "records.json", out)
    wb = openpyxl.load_workbook(out)
    assert "Orders" in wb.sheetnames
    assert "Warnings" in wb.sheetnames
    assert "Instructions" in wb.sheetnames


def test_export_row_count_matches_total_items(run_dir):
    import openpyxl
    out = run_dir / "orders_review.xlsx"
    export_review(run_dir / "records.json", out)
    wb = openpyxl.load_workbook(out)
    ws = wb["Orders"]
    data_rows = ws.max_row - 1
    total_items = sum(len(r["items"]) for r in RECORDS_DICT["records"])
    assert data_rows == total_items


def test_warning_rows_appear_on_warnings_sheet(run_dir):
    import openpyxl
    out = run_dir / "orders_review.xlsx"
    export_review(run_dir / "records.json", out)
    ws = openpyxl.load_workbook(out)["Warnings"]
    assert ws.max_row == 2


def test_flatten_produces_one_row_per_item():
    main, warnings = _flatten_records(RECORDS_DICT)
    assert len(main) == 3
    assert len(warnings) == 1


def test_flatten_default_include_is_Y():
    main, _ = _flatten_records(RECORDS_DICT)
    assert all(r["include"] == "Y" for r in main)


def _write_xlsx_edit(run_dir, edits: dict):
    import openpyxl

    out = run_dir / "orders_review.xlsx"
    export_review(run_dir / "records.json", out)

    wb = openpyxl.load_workbook(out)
    ws = wb["Orders"]
    headers = [c.value for c in ws[1]]

    for row in ws.iter_rows(min_row=2):
        row_dict = {headers[i]: (row[i].value or "") for i in range(len(headers))}
        key = (str(row_dict.get("Order ID", "")), str(row_dict.get("Product / Region", "")))
        if key in edits:
            for col_header, new_val in edits[key].items():
                col_idx = headers.index(col_header)
                row[col_idx].value = new_val

    wb.save(out)
    return out


def test_no_edit_produces_no_changes(run_dir):
    out = run_dir / "orders_review.xlsx"
    export_review(run_dir / "records.json", out)
    diff, _ = diff_review(run_dir / "records.json", out)
    assert not diff.has_changes


def test_student_name_edit_is_detected(run_dir):
    out = _write_xlsx_edit(run_dir, {
        ("5001", "large_labels"): {"Student Name": "Ahmed AlRashidi"},
    })
    diff, new_recs = diff_review(run_dir / "records.json", out)
    assert any("student_name" in c for c in diff.changed_fields)
    ahmed = next(r for r in new_recs["records"] if r["order_id"] == "5001")
    assert ahmed["student_name"] == "Ahmed AlRashidi"


def test_qty_edit_is_detected_and_applied(run_dir):
    out = _write_xlsx_edit(run_dir, {
        ("5001", "large_labels"): {"Qty": 3},
    })
    diff, new_recs = diff_review(run_dir / "records.json", out)
    assert any("qty" in c for c in diff.changed_fields)
    rec = next(r for r in new_recs["records"] if r["order_id"] == "5001")
    item = next(i for i in rec["items"] if i["region_id"] == "large_labels")
    assert item["qty"] == 3


def test_include_N_excludes_the_whole_order(run_dir):
    out = _write_xlsx_edit(run_dir, {
        ("5002", "large_labels"): {"Include?": "N"},
    })
    diff, new_recs = diff_review(run_dir / "records.json", out)
    assert len(diff.excluded_orders) == 1
    ids = [r["order_id"] for r in new_recs["records"]]
    assert "5002" not in ids and "5001" in ids


def test_invalid_qty_is_a_validation_error(run_dir):
    out = _write_xlsx_edit(run_dir, {
        ("5001", "large_labels"): {"Qty": "banana"},
    })
    diff, _ = diff_review(run_dir / "records.json", out)
    assert diff.validation_errors
    assert not diff.is_clean


def test_apply_review_writes_atomically(run_dir):
    out = _write_xlsx_edit(run_dir, {
        ("5001", "large_labels"): {"Student Name": "New Name"},
    })
    diff, new_recs = diff_review(run_dir / "records.json", out)
    records_path = run_dir / "records.json"
    apply_review(new_recs, records_path)

    assert not (run_dir / "records.json.tmp").exists()
    written = json.loads(records_path.read_text(encoding="utf-8"))
    ahmed = next(r for r in written["records"] if r["order_id"] == "5001")
    assert ahmed["student_name"] == "New Name"


def test_excluded_order_ids_recorded_in_output(run_dir):
    out = _write_xlsx_edit(run_dir, {
        ("5002", "large_labels"): {"Include?": "N"},
    })
    _, new_recs = diff_review(run_dir / "records.json", out)
    assert "5002" in new_recs["excluded_order_ids"]


def test_no_tag_run_is_complete_and_flagged(tmp_path):
    from orchestrator.journal import Journal

    d = tmp_path / "run"
    d.mkdir()
    j = Journal(d, d.name)
    j.record("export_begin", sheet_numbers=[1], order_ids=["5001"])
    j.record("sheet_written", sheet_number=1)
    j.record("report_written")
    j.record("tag_begin", order_ids=["5001"])
    j.record("export_complete", tagging_skipped=True)

    assert j.is_complete()
    assert j.tagging_was_skipped()
    assert j.untagged_order_ids() == ["5001"]


def test_normal_run_is_not_flagged_as_skipped(tmp_path):
    from orchestrator.journal import Journal

    d = tmp_path / "run"
    d.mkdir()
    j = Journal(d, d.name)
    j.record("export_begin", sheet_numbers=[1], order_ids=["5001"])
    j.record("tag_ok", order_id="5001")
    j.record("export_complete", tagging_skipped=False)

    assert j.is_complete()
    assert not j.tagging_was_skipped()
    assert j.untagged_order_ids() == []
